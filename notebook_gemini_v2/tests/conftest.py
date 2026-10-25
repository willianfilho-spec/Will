import io
import sys
from pathlib import Path

import pytest
from docx import Document
from reportlab.pdfgen import canvas

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from nb import repo  # noqa: E402
from nb.db import connect  # noqa: E402


def make_pdf(pages: list[str | None]) -> bytes:
    """Cada item vira uma página; None = página só com desenho (sem texto, como um scan)."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for text in pages:
        if text is None:
            c.rect(50, 50, 300, 300, fill=1)
        else:
            y = 800
            for line in text.split("\n"):
                c.drawString(72, y, line)
                y -= 16
        c.showPage()
    c.save()
    return buf.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None, heading: str | None = None) -> bytes:
    doc = Document()
    if heading:
        doc.add_heading(heading, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, value in enumerate(row):
                t.cell(i, j).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def conn(tmp_path):
    c = connect(tmp_path / "test.db")
    yield c
    c.close()


@pytest.fixture
def nid(conn):
    return repo.create_notebook(conn, "Notebook de teste")


@pytest.fixture
def docs_dir(tmp_path):
    return tmp_path / "documents"


def make_shadow_pdf(text: str, copies: int = 3) -> bytes:
    """Cada letra desenhada várias vezes quase no mesmo lugar (sombra/contorno de título)."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.setFont("Helvetica-Bold", 20)
    x = 72
    for ch in text:
        for k in range(copies):
            c.drawString(x + k * 0.4, 700, ch)
        x += c.stringWidth(ch, "Helvetica-Bold", 20)
    c.showPage()
    c.save()
    return buf.getvalue()


def make_pptx(slides: list[list[str] | None], notes: dict[int, str] | None = None) -> bytes:
    """Cada item é a lista de linhas de um slide; None = slide só com forma (sem texto)."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for i, lines in enumerate(slides, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        if lines is None:
            slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(2), Inches(2))
        else:
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
            box.text_frame.text = lines[0]
            for line in lines[1:]:
                box.text_frame.add_paragraph().text = line
        if notes and i in notes:
            slide.notes_slide.notes_text_frame.text = notes[i]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
