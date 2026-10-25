"""Extração de texto preservando a numeração das páginas.

"Página" aqui é a posição física no PDF (1 = primeira folha do arquivo), que pode
diferir do número impresso na folha. Formatos sem paginação própria (DOCX, TXT,
MD, CSV) viram uma única unidade com page_number = None.
"""
import io
import re
import unicodedata
from dataclasses import dataclass, field

from docx import Document
from pptx import Presentation
from docx.table import Table
from docx.text.paragraph import Paragraph
import pypdfium2 as pdfium
from pypdf import PdfReader
from pypdf.errors import PyPdfError

# Abaixo disso a página é tratada como "sem texto" (provável imagem escaneada).
MIN_CHARS_PER_PAGE = 20


class ExtractionError(Exception):
    """Falha esperada de leitura; a mensagem pode ser exibida ao usuário."""


@dataclass
class ExtractedPage:
    number: int | None
    text: str

    @property
    def has_text(self) -> bool:
        return len(self.text.strip()) >= MIN_CHARS_PER_PAGE


@dataclass
class ExtractedDocument:
    kind: str
    paginated: bool
    pages: list[ExtractedPage]
    warnings: list[str] = field(default_factory=list)

    @property
    def empty_pages(self) -> list[int]:
        return [p.number for p in self.pages if p.number is not None and not p.has_text]


# Ligaturas tipográficas (U+FB00..FB06) viram letras comuns, senão "diﬁculdade" não casa com
# "dificuldade" na busca. NFKC não serve: converteria também "º" em "o" (Art. 2º -> Art. 2o).
_LIGATURES = str.maketrans({
    "\ufb00": "ff", "\ufb01": "fi", "\ufb02": "fl", "\ufb03": "ffi", "\ufb04": "ffl",
    "\ufb05": "st", "\ufb06": "st",
})
# Alguns PDFs gravam o marcador de lista como DEL (0x7F); o pdfium lê isso como "•".
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b-\x1f\x7f-\x9f\ufffe\uffff]")


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x7f", "•")
    text = _CONTROL_CHARS.sub("", text).translate(_LIGATURES)
    # Alguns PDFs separam palavras com TAB em vez de espaço.
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r" ?\n ?", "\n", text)
    return unicodedata.normalize("NFC", text).strip()


def _looks_garbled(text: str) -> bool:
    """Texto com glifos sobrepostos (sombra/contorno) sai do pypdf como 'P P PO O ON N NT...' ou uma
    letra por linha. Assinatura: a mesma letra isolada repetida em sequência ('P P', 'O O'), o que
    não ocorre em português normal; 'a o e' sozinhos existem, mas nunca duas vezes seguidas."""
    tokens = text.split()
    repeated = sum(1 for a, b in zip(tokens, tokens[1:]) if len(a) == 1 and a.isalpha() and a == b)
    return repeated >= 3


def _pdfium_pages(data: bytes, numbers: list[int]) -> dict[int, str]:
    """Texto dessas páginas (1-based) pelo pdfium; vazio se ele não conseguir abrir o arquivo."""
    try:
        pdf = pdfium.PdfDocument(data)
        return {n: _clean(pdf[n - 1].get_textpage().get_text_range()) for n in numbers}
    except Exception:  # o pdfium levanta tipos próprios; aqui qualquer falha significa "sem alternativa"
        return {}


def extract_pdf(data: bytes) -> ExtractedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            if not reader.decrypt(""):
                raise ExtractionError("PDF protegido por senha. Remova a proteção e envie novamente.")
        raw_pages = [page.extract_text() or "" for page in reader.pages]
    except ExtractionError:
        raise
    except (PyPdfError, ValueError, KeyError, OSError, RecursionError) as exc:
        raise ExtractionError(f"PDF inválido ou corrompido ({type(exc).__name__}).") from exc
    if not raw_pages:
        raise ExtractionError("O PDF não tem páginas.")

    pages = [ExtractedPage(i, _clean(t)) for i, t in enumerate(raw_pages, start=1)]
    doc = ExtractedDocument("pdf", True, pages)

    garbled = [p.number for p in pages if _looks_garbled(p.text)]
    if garbled:
        alt = _pdfium_pages(data, garbled)
        fixed = [n for n in garbled if alt.get(n) and not _looks_garbled(alt[n])]
        for n in fixed:
            pages[n - 1].text = alt[n]
        if fixed:
            doc.warnings.append(
                f"Página(s) {_format_pages(fixed)}: o texto original vinha com letras repetidas "
                "(efeito de sombra/contorno) e foi relido por um segundo extrator. Confira na pré-visualização."
            )
        unfixed = [n for n in garbled if n not in fixed]
        if unfixed:
            doc.warnings.append(
                f"Página(s) {_format_pages(unfixed)}: texto possivelmente embaralhado e sem alternativa de leitura."
            )
    empty = doc.empty_pages
    if len(empty) == len(pages):
        raise ExtractionError(
            "Nenhum texto extraível. Provavelmente é um PDF digitalizado (imagem). "
            "Aplique OCR com uma ferramenta aprovada e envie a versão com texto."
        )
    if empty:
        doc.warnings.append(
            f"{len(empty)} de {len(pages)} página(s) sem texto extraível: {_format_pages(empty)}. "
            "Podem ser imagens que precisam de OCR; o conteúdo delas não será consultado."
        )
    return doc


def _iter_docx_blocks(document):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def extract_docx(data: bytes) -> ExtractedDocument:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx levanta tipos variados (zipfile, KeyError, lxml)
        raise ExtractionError(f"DOCX inválido ou corrompido ({type(exc).__name__}).") from exc

    lines: list[str] = []
    for block in _iter_docx_blocks(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            style = (block.style.name if block.style is not None else "") or ""
            lines.append(f"# {text}" if style.lower().startswith(("heading", "título")) else text)
        else:
            for row in block.rows:
                cells = []
                for cell in row.cells:
                    value = " ".join(cell.text.split())
                    if not cells or cells[-1] != value:  # células mescladas repetem o valor
                        cells.append(value)
                if any(cells):
                    lines.append(" | ".join(cells))
    text = _clean("\n".join(lines))
    if len(text) < MIN_CHARS_PER_PAGE:
        raise ExtractionError("Nenhum texto extraível no DOCX.")
    doc = ExtractedDocument("docx", False, [ExtractedPage(None, text)])
    doc.warnings.append("DOCX não tem páginas fixas; o texto será citado sem número de página.")
    return doc


def _pptx_shape_lines(shape) -> list[str]:
    lines: list[str] = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        lines += [" ".join(p.text.split()) for p in shape.text_frame.paragraphs if p.text.strip()]
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [" ".join(c.text.split()) for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    for child in getattr(shape, "shapes", []):  # grupos de formas
        lines += _pptx_shape_lines(child)
    return lines


def extract_pptx(data: bytes) -> ExtractedDocument:
    """Cada slide vira uma página (numeração = ordem dos slides)."""
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:  # python-pptx levanta tipos variados (zipfile, KeyError, lxml)
        raise ExtractionError(f"PPTX inválido ou corrompido ({type(exc).__name__}).") from exc
    pages: list[ExtractedPage] = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            lines += _pptx_shape_lines(shape)
        if slide.has_notes_slide:
            notes = " ".join(slide.notes_slide.notes_text_frame.text.split())
            if notes:
                lines.append(f"Notas do apresentador: {notes}")
        pages.append(ExtractedPage(number, _clean("\n".join(lines))))
    if not pages:
        raise ExtractionError("A apresentação não tem slides.")
    doc = ExtractedDocument("pptx", True, pages)
    empty = doc.empty_pages
    if len(empty) == len(pages):
        raise ExtractionError("Nenhum texto extraível nos slides (provavelmente só imagens). Exporte para PDF com texto ou aplique OCR.")
    if empty:
        doc.warnings.append(
            f"{len(empty)} de {len(pages)} slide(s) sem texto: {_format_pages(empty)}. "
            "Texto dentro de imagens não é lido."
        )
    return doc


def extract_text_file(data: bytes, kind: str) -> ExtractedDocument:
    warnings: list[str] = []
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
        warnings.append("Arquivo não estava em UTF-8; lido como Windows-1252. Confira acentos.")
    text = _clean(text)
    if len(text) < MIN_CHARS_PER_PAGE:
        raise ExtractionError("Arquivo vazio ou com texto insuficiente.")
    return ExtractedDocument(kind, False, [ExtractedPage(None, text)], warnings)


def extract(data: bytes, extension: str) -> ExtractedDocument:
    ext = extension.lower()
    if ext == ".pdf":
        return extract_pdf(data)
    if ext == ".docx":
        return extract_docx(data)
    if ext == ".pptx":
        return extract_pptx(data)
    if ext in (".txt", ".md", ".csv"):
        return extract_text_file(data, ext.lstrip("."))
    raise ExtractionError("Formato não suportado. Use PDF, DOCX, PPTX, TXT, MD ou CSV.")


def _format_pages(numbers: list[int]) -> str:
    """[1,2,3,7] -> '1-3, 7'"""
    ranges, start, prev = [], None, None
    for n in numbers:
        if start is None:
            start = prev = n
        elif n == prev + 1:
            prev = n
        else:
            ranges.append((start, prev))
            start = prev = n
    if start is not None:
        ranges.append((start, prev))
    return ", ".join(str(a) if a == b else f"{a}-{b}" for a, b in ranges)
