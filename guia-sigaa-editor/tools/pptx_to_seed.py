#!/usr/bin/env python3
"""Converte o PPTX do Guia SIGAA em fichas editáveis (seed.json + imagens WebP).

Uso: python3 pptx_to_seed.py guia.pptx saida_dir
O conversor lê posição, tamanho de fonte e negrito de cada caixa de texto e
reconhece os padrões do guia (caminho, passos, quadros, cartões, fluxo, telas).
Todo texto que não encaixa em um padrão vira bloco de texto, para nada se perder.
"""
import base64, io, json, re, sys
from pptx import Presentation
from PIL import Image

EMU_X, EMU_Y = 7559675 / 210.0, 10691813 / 297.0
QUADRO_VAR = {
    'ATENÇÃO': 'atencao', 'REGRA CRÍTICA': 'atencao',
    'OBSERVAÇÃO': 'observacao', 'OBERVAÇÃO': 'observacao',
    'QUANDO USAR': 'quando',
    'REGULAMENTAÇÃO': 'regulamentacao', 'NORMATIVAS RELACIONADAS': 'regulamentacao',
}
LABELS_FIXOS = {'COMO FAZER', 'TELA DO SIGAA', 'TELAS DO SIGAA', 'CAMINHO NO SIGAA'}


def webp(blob, maxw, q):
    im = Image.open(io.BytesIO(blob))
    if im.width > maxw:
        im = im.resize((maxw, round(im.height * maxw / im.width)), Image.LANCZOS)
    if im.mode not in ('RGB', 'RGBA'):
        im = im.convert('RGBA' if 'transparency' in im.info else 'RGB')
    out = io.BytesIO()
    im.save(out, 'WEBP', quality=q, method=6)
    return 'data:image/webp;base64,' + base64.b64encode(out.getvalue()).decode()


class It:
    def __init__(self, sh):
        self.x, self.y = sh.left / EMU_X, sh.top / EMU_Y
        self.w, self.h = sh.width / EMU_X, sh.height / EMU_Y
        self.pic = sh.shape_type == 13
        self.used = False
        if self.pic:
            self.sh = sh
            self.text = ''
            return
        paras = [p for p in sh.text_frame.paragraphs]
        self.lines = []
        for p in paras:
            t = ''.join(r.text for r in p.runs).replace('\x0b', ' ').replace('​', '').strip()
            if t:
                self.lines.append(re.sub(r'[ \t]+', ' ', t))
        self.text = '\n'.join(self.lines)
        run = next((r for p in paras for r in p.runs if r.text.strip()), None)
        self.size = run.font.size.pt if run is not None and run.font.size else 0
        self.bold = bool(run.font.bold) if run is not None else False
        self.bottom = self.y + self.h

    @property
    def upper(self):
        t = self.text.strip()
        return bool(t) and t == t.upper() and any(c.isalpha() for c in t)

    @property
    def mostly_upper(self):
        letras = [c for c in self.text if c.isalpha()]
        return len(letras) >= 3 and sum(c.isupper() for c in letras) / len(letras) >= 0.6

    def __repr__(self):
        return f'<{self.x:.0f},{self.y:.0f} {self.text[:30]!r}>'


def is_chrome(it):
    if it.pic:
        return it.x > 165 and it.y < 25 and it.w < 25
    t = it.text.strip()
    if not t:
        return True
    if it.y < 8 or it.y >= 288:
        return True
    if t == 'GUIA SIGAA' and it.size == 9:
        return True
    if t.startswith('DIOAC | PROEN | UFMA'):
        return True
    if t.isdigit() and it.x >= 165 and it.y >= 200:
        return True
    return False


def quadro(rot, texto, y, x, w, h):
    var = QUADRO_VAR.get(rot.strip().upper(), 'observacao')
    b = {'tipo': 'quadro', 'variante': var, 'rotulo': rot.strip(), 'texto': texto}
    return dict(b=b, y0=y, y1=y + h, x=x, w=w)


def blocks10(items):
    t = {i.text.strip(): i for i in items}
    def g(k):
        return next(i for i in items if i.text.strip().startswith(k))
    cards = []
    for n, tit, corpo in (('01', 'SOLICITAR', 'Novo plano'), ('02', 'ANALISAR', 'Coordenação pode'),
                          ('03', 'HOMOLOGAR PAGAMENTO', 'Após o deferimento'), ('04', 'EXECUTAR', 'Aluno retorna')):
        cards.append(f'{n}  {g(tit).text.replace(chr(10), " ")}\n{g(corpo).text}')
    for i in items:
        i.used = True
    tab = 'ETAPA;SE DEFERIDO;SE INDEFERIDO\n' + \
        'ANÁLISE DA COORDENAÇÃO;HOMOLOGAÇÃO DE PAGAMENTO/ISENÇÃO;DESLIGAMENTO DO ALUNO\n' + \
        'HOMOLOGAÇÃO DE PAGAMENTO/ISENÇÃO;PLANO EM EXECUÇÃO;INCLUSÃO DE NOVO COMPROVANTE'
    return [
        dict(b={'tipo': 'fluxo', 'itens': '\n\n'.join(cards), 'colunas': 4}, y0=60, y1=100, x=17, w=190),
        dict(b={'tipo': 'secao', 'texto': 'STATUS E AÇÕES PRINCIPAIS'}, y0=110, y1=115, x=20, w=100),
        dict(b={'tipo': 'tabela', 'csv': tab}, y0=118, y1=160, x=20, w=190),
    ]


def tabela24(items):
    hdr = [i for i in items if i.text.strip() in ('QUEM?', 'OPERAÇÃO', 'ITEM DO SUBMENU')]
    if len(hdr) < 3:
        return []
    y0 = hdr[0].y - 1
    fim = next((i.y for i in items if i.text.startswith('DEFINIÇÃO')), 1e9)
    cols = sorted(hdr, key=lambda i: i.x)
    cells = [i for i in items if not i.used and y0 <= i.y < fim - 2 and i.text.strip() != 'CAMINHO NO SIGAA']
    cells = [i for i in cells if i not in hdr and any(abs(i.x - c.x) < 6 for c in cols)]
    rows = {}
    first = sorted([i for i in cells if abs(i.x - cols[0].x) < 6], key=lambda i: i.y)
    for i in cells:
        r = min(first, key=lambda f: abs(f.y - i.y))
        c = min(range(3), key=lambda k: abs(i.x - cols[k].x))
        rows.setdefault(id(r), [r.y, ['', '', '']])[1][c] = i.text.replace('\n', ' ').replace(';', ',')
        i.used = True
    for h in hdr:
        h.used = True
    lines = [';'.join(h.text for h in cols)] + [';'.join(v[1]) for v in sorted(rows.values(), key=lambda v: v[0])]
    return [dict(b={'tipo': 'tabela', 'csv': '\n'.join(lines)}, y0=y0, y1=fim - 2, x=cols[0].x, w=190)]


def convert_segment(items, pics, secao, titulo, num_slide):
    """items: textos do segmento já sem título/cabeçalho; devolve lista de blocos posicionados."""
    blocks = []
    if num_slide == 10:
        return blocks10(items)
    if num_slide == 24:
        blocks += tabela24(items)
    # --- fontes de imagem
    fontes = [i for i in items if i.text.startswith('Fonte:')]
    for f in fontes:
        f.used = True
    # --- caminho
    for lab in [i for i in items if i.text.strip() == 'CAMINHO NO SIGAA' and not i.used]:
        cands = [i for i in items if not i.used and i is not lab and i.x > lab.x + 25 and abs(i.y - lab.y) < 14]
        lab.used = True
        if cands:
            c = min(cands, key=lambda i: abs(i.y - lab.y))
            c.used = True
            blocks.append(dict(b={'tipo': 'caminho', 'texto': c.text}, y0=min(lab.y, c.y), y1=max(lab.bottom, c.bottom), x=lab.x, w=190))
    # --- tipos de linha única "ROTULO  texto" (regulamentação inline)
    for i in items:
        if not i.used and i.text.startswith('REGULAMENTAÇÃO') and len(i.text) > 20 and i.w > 150:
            body = i.text[len('REGULAMENTAÇÃO'):].strip()
            blocks.append(quadro('REGULAMENTAÇÃO', body, i.y, i.x, i.w, i.h))
            i.used = True
    # --- cartões de pessoas (duas linhas: NOME COMPLETO / função)
    pess = [i for i in items if not i.used and i.text.startswith('NOME COMPLETO') and '\n' in i.text]
    if pess:
        pess.sort(key=lambda i: (round(i.y / 6), i.x))
        for i in pess: i.used = True
        ncol = len({round(i.x / 25) for i in pess})
        blocks.append(dict(b={'tipo': 'cartoes', 'itens': '\n\n'.join(i.text for i in pess), 'colunas': ncol},
                           y0=min(i.y for i in pess), y1=max(i.bottom for i in pess), x=pess[0].x, w=190, row=True))
    # --- pares rótulo + corpo
    heads = [i for i in items if not i.used and i.bold and 9 <= i.size <= 13.5 and i.mostly_upper and '\n' not in i.text and i.text.strip() not in LABELS_FIXOS
             and not i.text.strip().isdigit() and i.text.strip() != '>']
    pairs = []
    for a in sorted(heads, key=lambda i: (i.y, i.x)):
        bs = [b for b in items if not b.used and b is not a and b.y > a.y and b.y - a.y <= 17 and abs(b.x - a.x) <= 7 and (a.w < 100 or b.w >= 100)
              and not (b.bold and 9 <= b.size <= 13.5 and b.mostly_upper and '\n' not in b.text) and not b.text.strip().isdigit() and b.text.strip() != '>']
        if bs:
            b = min(bs, key=lambda i: i.y)
            if any(h is not a and not h.used and a.y < h.y < b.y and abs(h.x - b.x) <= 7 for h in heads):
                continue
            a.used = b.used = True
            pairs.append((a, b))
    rows = []
    for a, b in sorted(pairs, key=lambda p: (p[0].y, p[0].x)):
        for r in rows:
            if abs(r[0][0].y - a.y) <= 4:
                r.append((a, b)); break
        else:
            rows.append([(a, b)])
    for r in rows:
        r.sort(key=lambda p: p[0].x)
        if len(r) == 1:
            a, b = r[0]
            blocks.append(quadro(a.text.replace('\n', ' '), b.text, a.y, a.x, max(a.w, b.w), b.bottom - a.y))
        else:
            fluxo = all(re.match(r'^\d\d\b', a.text) for a, b in r)
            txt = '\n\n'.join(a.text.replace('\n', ' ') + '\n' + b.text for a, b in r)
            blocks.append(dict(b={'tipo': 'fluxo' if fluxo else 'cartoes', 'itens': txt, 'colunas': len(r)},
                               y0=r[0][0].y, y1=max(b.bottom for a, b in r), x=r[0][0].x, w=190, row=True))
    # --- passos
    circles = [i for i in items if not i.used and i.text.strip().isdigit() and i.w <= 10.5 and i.h >= 6.3]
    markers = [i for i in items if not i.used and i.text.strip().isdigit() and i.w <= 10.5 and i.h < 6.3]
    for m in markers:
        m.used = True  # números soltos de legenda sobre os prints
    steps = []
    for c in circles:
        c.used = True
        steps.append({'n': int(c.text), 'c': c, 'its': []})
    others = [i for i in items if not i.used and i.text.strip() != '>']
    for i in others:
        best, score = None, 1e9
        for s in steps:
            c = s['c']
            dx, dy = i.x - c.x, i.y - c.y
            if dx < 8 or dx > 45 or dy < -5:
                continue
            sc = abs(dx - 13) + dy
            if sc < score:
                best, score = s, sc
        if best is not None and score < 80:
            best['its'].append(i); i.used = True
    # rótulos
    rotulo_passos = 'COMO FAZER'
    for i in items:
        if i.text.strip() == 'COMO FAZER':
            i.used = True
    tela_labels = [i for i in items if i.text.strip() in ('TELA DO SIGAA', 'TELAS DO SIGAA')]
    # rótulo de imagem: o que fica logo acima de uma imagem
    tag_for = {}
    for lab in tela_labels:
        near = [p for p in pics if 0 <= p['y'] - lab.y <= 14]
        if near:
            tag_for[id(min(near, key=lambda p: p['y']))] = lab.text.strip()
            lab.used = True
        elif any(0 <= s['c'].y - lab.y <= 12 for s in steps):
            rotulo_passos = lab.text.strip(); lab.used = True
    steps.sort(key=lambda s: (round(s['c'].y / 9), s['c'].x))
    # agrupar passos consecutivos sem imagem no meio
    grupos = []
    for s in steps:
        y = s['c'].y
        if grupos:
            last = grupos[-1][-1]
            if not any(last['c'].y < p['y'] < y for p in pics) and s['n'] == last['n'] + 1 and y - last['c'].y < 60:
                grupos[-1].append(s); continue
        grupos.append([s])
    for gi, g in enumerate(grupos):
        itens = []
        for s in g:
            its = sorted(s['its'], key=lambda i: (i.y, i.x))
            ls = [i.text for i in its]
            itens.append('\n'.join(ls))
        xs = sorted({round(s['c'].x / 20) for s in g})
        cols = 1 if len(xs) == 1 else len(xs)
        ys = [s['c'].y for s in g] + [i.bottom for s in g for i in s['its']]
        b = {'tipo': 'passos', 'rotulo': rotulo_passos if (gi == 0 or g[0]['n'] == 1) else '', 'itens': '\n\n'.join(itens), 'inicio': g[0]['n'], 'colunas': cols}
        blocks.append(dict(b=b, y0=min(ys) - (6 if gi == 0 else 0), y1=max(ys), x=min(s['c'].x for s in g), w=190 if cols > 1 else 90, passos=True))
    # --- imagens
    for p in pics:
        fonte = ''
        cand = [f for f in fontes if -3 <= f.y - (p['y'] + p['h']) <= 6]
        if cand:
            fonte = cand[0].text.replace('Fonte: ', '').replace('Fonte:', '').strip()
            fonte = 'Fonte: ' + fonte
        b = {'tipo': 'imagem', 'src': p['key'], 'rotulo': tag_for.get(id(p), ''), 'legenda': '', 'fonte': fonte,
             'largura': max(20, min(100, round(p['w'] / 184 * 100)))}
        blocks.append(dict(b=b, y0=p['y'] - (5 if tag_for.get(id(p)) else 0), y1=p['y'] + p['h'], x=p['x'], w=p['w']))
    for f in fontes:
        if not any(-3 <= f.y - (p['y'] + p['h']) <= 6 for p in pics):
            blocks.append(dict(b={'tipo': 'texto', 'texto': f.text}, y0=f.y, y1=f.bottom, x=f.x, w=f.w))
    # --- sobras em negrito com título + corpo na mesma caixa (cartões)
    duplas = [i for i in items if not i.used and i.bold and '\n' in i.text and i.w < 100]
    if len(duplas) >= 2:
        duplas.sort(key=lambda i: (round(i.y / 8), i.x))
        for i in duplas: i.used = True
        blocks.append(dict(b={'tipo': 'cartoes', 'itens': '\n\n'.join(i.text for i in duplas), 'colunas': 2},
                           y0=min(i.y for i in duplas), y1=max(i.bottom for i in duplas), x=duplas[0].x, w=190, row=True))
    # --- rótulo solto que fica logo acima de um quadro (ex.: QUANDO USAR / OBSERVAÇÃO)
    for i in items:
        if i.used or i.text.strip() not in QUADRO_VAR:
            continue
        alvo = [d for d in blocks if d['b']['tipo'] == 'quadro' and 0 <= d['y0'] - i.y <= 16 and abs(d['x'] - i.x) <= 12]
        if alvo:
            d = min(alvo, key=lambda d: d['y0'])
            d['b']['rotulo'] = i.text.strip() + ' · ' + d['b']['rotulo']
            d['y0'] = min(d['y0'], i.y)
            i.used = True
    # --- sobras: cabeçalhos (secao) e textos
    for i in sorted([i for i in items if not i.used and i.text.strip() != '>' and i.text.strip() not in ('TELA DO SIGAA', 'TELAS DO SIGAA', 'COMO FAZER')], key=lambda i: (i.y, i.x)):
        if i.upper and len(i.text) <= 70 and '\n' not in i.text:
            blocks.append(dict(b={'tipo': 'secao', 'texto': i.text}, y0=i.y, y1=i.bottom, x=i.x, w=i.w))
        else:
            blocks.append(dict(b={'tipo': 'texto', 'texto': i.text}, y0=i.y, y1=i.bottom, x=i.x, w=i.w))
        i.used = True
    for i in items:
        if not i.used and i.text.strip() not in ('>', 'TELA DO SIGAA', 'TELAS DO SIGAA', 'COMO FAZER'):
            print('  [aviso] texto não aproveitado:', i)
    return blocks


def order_blocks(blocks):
    """Ordena por posição vertical e coloca lado a lado (meia largura) o que estava lado a lado."""
    blocks.sort(key=lambda d: (d['y0'], d['x']))
    out, used = [], set()
    for k, d in enumerate(blocks):
        if k in used:
            continue
        out.append(d)
        used.add(k)
        if d['w'] < 100 and d['b']['tipo'] in ('quadro', 'passos', 'imagem'):
            left = d['x'] < 100
            for j, e in enumerate(blocks):
                if j in used or j == k or e['w'] >= 100 or e['b']['tipo'] not in ('quadro', 'passos'):
                    continue
                if (e['x'] < 100) == left:
                    continue
                if e['y0'] <= d['y1'] + 2 and d['y0'] <= e['y1'] + 2:
                    used.add(j)
                    pair = [d, e] if left else [e, d]
                    out.pop()
                    out.extend(pair)
                    for q in pair:
                        q['b']['meia'] = True
                    break
    return out


def merge_cards(blocks):
    res = []
    for d in blocks:
        if res and d['b']['tipo'] in ('cartoes', 'fluxo') and res[-1]['b']['tipo'] == d['b']['tipo'] \
                and res[-1]['b']['colunas'] == d['b']['colunas'] and d['y0'] - res[-1]['y1'] < 45:
            res[-1]['b']['itens'] += '\n\n' + d['b']['itens']
            res[-1]['y1'] = d['y1']
        else:
            res.append(d)
    return res


def main(pptx_path, outdir):
    import os
    os.makedirs(outdir, exist_ok=True)
    prs = Presentation(pptx_path)
    images = {}
    fichas = []
    cfg = {}
    for sn, slide in enumerate(prs.slides, 1):
        its = [It(sh) for sh in slide.shapes if sh.shape_type == 13 or sh.has_text_frame]
        # logos
        for it in its:
            if it.pic and it.w > 40 and it.y < 120 and sn == 1:
                images['logo'] = webp(it.sh.image.blob, 760, 85)
            if it.pic and is_chrome(it) and 'logo_pequena' not in images:
                images['logo_pequena'] = webp(it.sh.image.blob, 360, 85)
        if sn == 1:
            t = [i for i in its if not i.pic and i.text.strip()]
            g = lambda s: next((i.text for i in t if i.text.strip().startswith(s)), '')
            fichas.append({'tipo': 'capa', 'nome': 'Capa', 'capa': {
                'titulo': g('GUIA SIGAA'), 'pilula': g('Coordenação'), 'sub': g('Orientações'), 'rodape': g('DIOAC')}, 'blocos': []})
            continue
        its = [i for i in its if not is_chrome(i)]
        texts = [i for i in its if not i.pic]
        pics = []
        for i in its:
            if i.pic:
                key = f'p{sn}_{len(pics)}'
                images[key] = webp(i.sh.image.blob, 1100, 74)
                pics.append(dict(y=i.y, x=i.x, w=i.w, h=i.h, key='seed:' + key))
        titles = sorted([i for i in texts if i.bold and i.size >= 22], key=lambda i: i.y)
        if not titles:
            print('slide', sn, 'sem título'); continue
        t0 = titles[0]
        lab = [i for i in texts if i.bold and i.size == 10 and i.upper and i.y < t0.y and t0.y - i.y < 20 and i.w > 60]
        secao = max(lab, key=lambda i: i.y).text if lab else ''
        if lab: max(lab, key=lambda i: i.y).used = True
        if sn == 3:
            secao = ''
        for t in titles: t.used = True
        for k, t in enumerate(titles):
            nxt = titles[k + 1].y - 0.5 if k + 1 < len(titles) else 1e9
            seg = [i for i in texts if not i.used and t.y <= i.y < nxt]
            sp = [p for p in pics if t.y <= p['y'] < nxt]
            blocks = convert_segment(seg, sp, secao, t.text, sn)
            blocks = merge_cards(order_blocks(blocks))
            if sn == 3:
                for d in blocks:
                    if d['b']['tipo'] == 'cartoes':
                        d['b']['estilo'] = 'topo'; d['b']['bullets'] = 'sim'
            bl = [{'id': f'b{sn}_{k}_{n}', 'meia': False, **d['b']} for n, d in enumerate(blocks)]
            bl.insert(0, {'id': f'b{sn}_{k}_t', 'tipo': 'titulo', 'meia': False, 'secao': secao, 'texto': t.text})
            fichas.append({'id': f'f{sn}_{k}', 'tipo': 'ficha', 'nome': t.text, 'origem': f'slide {sn}', 'blocos': bl})
    for n, f in enumerate(fichas):
        f.setdefault('id', f'f{n}')
    seed = {'fichas': fichas, 'images': images}
    json.dump(seed, open(os.path.join(outdir, 'seed.json'), 'w'), ensure_ascii=False)
    print('fichas', len(fichas), 'imagens', len(images), 'bytes', os.path.getsize(os.path.join(outdir, 'seed.json')))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
